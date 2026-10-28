"""Builds exam-style Excel files for every case in model.js.

For each case it writes two files into files/:
  <id>-blank-template.xlsx  historicals + assumptions filled, forecast and DCF cells blank (what the exam hands out)
  <id>-answer-key.xlsx same file with every formula filled in, matching the class model

The layout is the professor's Session 12 practice file (tools/base.xlsx), so rows and columns match class.
Run: <python with openpyxl> tools/make_templates.py   (needs node on PATH to read model.js)
"""
import json, subprocess, sys, copy
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
YEARS = ['FY27', 'FY28', 'FY29', 'FY30']
F = ['F', 'G', 'H', 'I']          # forecast columns on the P&L sheets
A = ['C', 'D', 'E', 'F']          # same years on Assumptions and DCF
JS = "const D=require('./model.js');const o={};for(const k in D.CASES){const c=D.CASES[k];o[k]={c,h:D.hist(c).inputs,v:D.run(c).v}};console.log(JSON.stringify(o))"
CASES = json.loads(subprocess.check_output(['node', '-e', JS], cwd=ROOT))


def forecast_formulas(c):
    """Every forecast-sheet formula, as in the class answer file. Returns {sheet: {cell: formula}}."""
    off, d2c, con = {}, {}, {}
    for i, (L, a) in enumerate(zip(F, A)):
        P = chr(ord(L) - 1)
        off.update({f'{L}6': f'={P}6*(1+Assumptions!{a}6)', f'{L}7': f'={P}7*(1+Assumptions!{a}7)', f'{L}8': f'={L}6+{L}7',
                    f'{L}9': f'={L}8*(1-Assumptions!{a}8)', f'{L}10': f'={L}8-{L}9', f'{L}11': f'={L}$8*Assumptions!{a}9',
                    f'{L}12': f'={L}$8*Assumptions!{a}10', f'{L}13': f'={L}10-{L}11-{L}12', f'{L}14': f'={L}$8*Assumptions!{a}11',
                    f'{L}15': f'={L}13-{L}14', f'{L}16': f'={P}16*(1+Assumptions!{a}12)', f'{L}17': f'={P}17*(1+Assumptions!{a}13)',
                    f'{L}18': f'={L}15-{L}16-{L}17'})
        d2c.update({f'{L}6': f'={P}6*(1+Assumptions!{a}17)', f'{L}7': f'={P}7*(1+Assumptions!{a}18)', f'{L}8': f'={L}6+{L}7',
                    f'{L}9': f'={L}8*(1-Assumptions!{a}19)', f'{L}10': f'={L}8-{L}9', f'{L}11': f'={L}$8*Assumptions!{a}20',
                    f'{L}12': f'={L}$8*Assumptions!{a}21', f'{L}13': f'={L}$8*Assumptions!{a}22', f'{L}14': f'={L}10-{L}11-{L}12-{L}13',
                    f'{L}15': f'={L}$8*Assumptions!{a}23', f'{L}16': f'={L}14-{L}15'})
        dep = '=Assumptions!$C$34+Assumptions!$C$38' if i >= c['capexYear'] else '=Assumptions!$C$34'
        con.update({f'{L}6': f"='Offline retail'!{L}8", f'{L}7': f'=D2C!{L}8', f'{L}8': f'={L}6+{L}7',
                    f'{L}9': f"='Offline retail'!{L}18", f'{L}10': f'=D2C!{L}16', f'{L}11': f'={L}9+{L}10',
                    f'{L}12': f'={P}12*(1+Assumptions!{a}27)', f'{L}13': f'={P}13*(1+Assumptions!{a}28)',
                    f'{L}14': f'={P}14*(1+Assumptions!{a}29)', f'{L}15': f'={P}15*(1+Assumptions!{a}30)',
                    f'{L}16': f'={L}11-SUM({L}12:{L}15)', f'{L}17': dep, f'{L}18': f'={L}16-{L}17',
                    f'{L}19': '=Assumptions!$C$42*Assumptions!$C$43', f'{L}20': f'={L}18-{L}19',
                    f'{L}21': f'=MAX(0,{L}20*Assumptions!$C$41)', f'{L}22': f'={L}20-{L}21'})
    return {'Offline retail': off, 'D2C': d2c, 'Consolidated ': con}


def dcf_formulas(c):
    f = {'C4': c['r'], 'C5': c['g'], 'F15': '=F14*(1+C5)/(C4-C5)', 'F21': '=F15*F18', 'C23': '=C20+D20+E20+F20+F21',
         'C27': '=C23', 'C28': c['ask'], 'C29': '=C27-C28', 'C33': '=-C28', 'C34': '=IRR(C33:G33)',
         'C38': 1.2, 'D38': '=D39*C38', 'C39': 1, 'D39': '=C23', 'C40': 0.8, 'D40': '=D39*C40',
         'C44': '=Assumptions!C40'}
    for i, (col, L) in enumerate(zip(A, F)):
        nxt = chr(ord(col) + 1)
        f.update({f'{col}9': f"='Consolidated '!{L}22", f'{col}10': f"='Consolidated '!{L}17", f'{col}11': f'={col}9+{col}10',
                  f'{col}12': f'={nxt}45', f'{col}13': c['capex'] if i == c['capexYear'] else 0,
                  f'{col}14': f'={col}11-{col}12-{col}13', f'{col}17': i + 1, f'{col}18': f'=1/(1+$C$4)^{col}17',
                  f'{col}20': f'={col}14*{col}18', f'{nxt}33': f'={col}14+{col}15',
                  f'{nxt}44': f"=Assumptions!$C$39*'Consolidated '!{L}8", f'{nxt}45': f'={nxt}44-{col}44'})
    return f


def fill_inputs(wb, case):
    c, h = case['c'], case['h']
    sheets = {'off': wb['Offline retail'], 'dc': wb['D2C'], 'con': wb['Consolidated ']}
    for key, ws in sheets.items():
        for row, vals in h[key].items():
            for col, v in zip('CDE', vals):
                ws[f'{col}{row}'] = v
    for col in 'CDE':
        sheets['con'][f'{col}19'] = 0
    asn = wb['Assumptions']
    o, d, k = c['off'], c['d2c'], c['common']
    rows = {6: o['gB'], 7: o['gA'], 8: o['gm'], 9: o['pack'], 10: o['pg'], 11: o['pm'], 12: o['gRent'], 13: o['gStaff'],
            17: d['gB'], 18: d['gA'], 19: d['gm'], 20: d['pack'], 21: d['pg'], 22: d['ship'], 23: d['pm'],
            27: k['gSal'], 28: k['gBrand'], 29: k['gTech'], 30: k['gOth']}
    for row, vals in rows.items():
        for col, v in zip(A, vals):
            asn[f'{col}{row}'] = v
    cy = YEARS[c['capexYear']]
    asn['B35'] = f"{c['asset'][0].upper() + c['asset'][1:]}, total cost (₹ cr)"
    asn['B38'] = f'Additional depreciation ({cy} onwards)'
    for cell, v in {'C34': c['dep0'], 'C35': c['capex'], 'C36': cy, 'C37': c['life'], 'C39': c['wcPct'], 'C40': c['wc0'],
                    'C41': c['tax'], 'C42': c['debt'], 'C43': c['intRate']}.items():
        asn[cell] = v
    asn['B44'] = None; asn['C44'] = None          # book equity is not used anywhere in the model
    dcf = wb['DCF (FCFE)']
    dcf['B7'] = f"Calculating value of {c['name']} using DCF"
    dcf['B13'] = f"Less: Capex for {c['asset']}"


def build(case_id, key):
    case = CASES[case_id]
    wb = openpyxl.load_workbook(ROOT / 'tools' / 'base.xlsx')
    fill_inputs(wb, case)
    forecast = forecast_formulas(case['c'])
    dcf = wb['DCF (FCFE)']
    yellow = copy.copy(wb['Offline retail']['C6'].fill)
    for sheet, cells in forecast.items():
        ws = wb[sheet]
        for cell, fml in cells.items():
            ws[cell] = fml if key else None
    for cell, val in dcf_formulas(case['c']).items():
        dcf[cell] = val if key else None
    if not key:   # mark where the exam expects typed inputs on the DCF sheet
        for cell in ['C4', 'C5', 'C28']:
            dcf[cell].fill = copy.copy(yellow)
    out = ROOT / 'files' / f"{case_id}-{'answer-key' if key else 'blank-template'}.xlsx"
    wb.save(out)
    return out


if __name__ == '__main__':
    (ROOT / 'files').mkdir(exist_ok=True)
    for cid in CASES:
        for key in (False, True):
            print(build(cid, key).name)
